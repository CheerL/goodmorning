import csv
import os

from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference


def find_csv_path(target_time_str, db_path):
    return [each for each in os.listdir(db_path) if target_time_str in each]

def add_sheet(wb, db_path, csv_path):
    symbol = csv_path.split('_')[1]
    sheetname = symbol[:-4]
    wb.create_sheet(sheetname)
    ws = wb[sheetname]
    with open(os.path.join(db_path, csv_path), 'r') as csv_file:
        reader = csv.reader(csv_file)
        for index, row in enumerate(reader):
            if index > 0:
                row = [float(each) for each in row]
                row[0] = round(row[0], 3)
            ws.append(row)

    increase_data = Reference(ws, range_string=f'{sheetname}!$D$1:$D${reader.line_num}')
    back_data = Reference(ws, range_string=f'{sheetname}!$F$1:$F${reader.line_num}')
    times = Reference(ws, range_string=f'{sheetname}!$A$2:$A${reader.line_num}')

    lc = LineChart()
    lc.title = sheetname
    lc.width = 30
    lc.height = 18
    lc.add_data(increase_data, titles_from_data=True)
    lc.add_data(back_data, titles_from_data=True)
    lc.set_categories(times)
    lc.x_axis.tickLblPos = 'low'
    lc.x_axis.title = '时间'
    lc.y_axis.title = '幅度'
    lc.y_axis.numFmt = '0.00%'
    ws.add_chart(lc, 'I3')

def create_excel(target_time_str, db_path):
    wb = Workbook()
    wb.remove(wb['Sheet'])
    for csv_path in find_csv_path(target_time_str, db_path):
        add_sheet(wb, db_path, csv_path)
    wb.save(os.path.join(db_path, '..', f'{target_time_str}.xlsx'))


