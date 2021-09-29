
import sys
import argparse
import os
import time
import datetime
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from utils import ROOT, config
from dataset.pgsql import get_session, get_Trade, Target

### Too old

DB_PATH = os.path.join(ROOT, 'test', 'db')
PGHOST = config.get('data', 'PGHost')

def find_csv_path(target_time_str, db_path):
    return [each for each in os.listdir(db_path) if target_time_str in each]

def add_sheet(wb, session, symbol, start, end):
    sheetname = symbol[:-4]
    wb.create_sheet(sheetname)
    ws = wb[sheetname]
    ws.append(['时间', '价格', '成交额', '涨幅', '高', '回撤'])
    open_ = 0
    high = 0
    vol = 0
    Trade = get_Trade(int(start))
    data = Trade.get_data(session, symbol, start, end).all()
    for index, trade in enumerate(data):
        time_ = round(float(trade.ts) / 1000 - start, 3)
        price = trade.price
        amount = trade.amount
        if index == 0:
            open_ = price
        increase = price / open_ - 1
        high = max(high, price)
        back = 1 - high / price
        vol += price * amount
        ws.append([time_, price, vol, increase, high, back])

    increase_data = Reference(ws, range_string=f'{sheetname}!$D$1:$D${len(data)+1}')
    back_data = Reference(ws, range_string=f'{sheetname}!$F$1:$F${len(data)+1}')
    times = Reference(ws, range_string=f'{sheetname}!$A$2:$A${len(data)+1}')

    lc = LineChart()
    lc.title = sheetname
    lc.width = 30
    lc.height = 18
    lc.add_data(increase_data, titles_from_data=True)
    lc.add_data(back_data, titles_from_data=True)
    lc.set_categories(times)
    lc.x_axis.tickLblPos = 'low'
    lc.x_axis.title = '时间'
    lc.y_axis.title = '幅度'
    lc.y_axis.numFmt = '0.00%'
    ws.add_chart(lc, 'I3')

def create_excel(target_time_str, db_path, targets, host):
    target_time = time.strptime(target_time_str, '%Y-%m-%d %H:%M:%S')
    target_time_ts = time.mktime(target_time)
    target_tm = time.strftime('%Y-%m-%d-%H', target_time)
    
    wb = Workbook()
    wb.remove(wb['Sheet'])
    with get_session(host=host) as session:
        if not targets:
            targets = session.query(Target).filter(Target.tm == target_tm).first().targets

        targets = targets.split(',')
        print(targets)
        for symbol in targets:
            add_sheet(wb, session, symbol, target_time_ts, target_time_ts + 300)
    wb.save(os.path.join(db_path, f'{target_tm}.xlsx'))


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('-t', '--time', default=datetime.date.today().strftime('%Y-%m-%d %H:%M:%S'))
    parser.add_argument('-s', '--symbol', default='')
    parser.add_argument('-H', '--host', default=PGHOST)
    parser.add_argument('-p', '--path', default=DB_PATH)
    args = parser.parse_args()

    create_excel(args.time, args.path, args.symbol, args.host)

if __name__ == '__main__':
    main()